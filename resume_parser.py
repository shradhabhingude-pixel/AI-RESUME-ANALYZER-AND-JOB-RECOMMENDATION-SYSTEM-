"""In-memory text extraction for PDF and DOCX resumes."""
from io import BytesIO
from pathlib import Path
from docx import Document
from pypdf import PdfReader

def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Extract PDF pages or DOCX paragraphs/tables without writing the resume to disk."""
    suffix = Path(filename).suffix.casefold()
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        if reader.is_encrypted:
            raise ValueError("This PDF is password-protected. Upload an unlocked copy.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = Document(BytesIO(file_bytes))
        texts = [p.text for p in doc.paragraphs]
        texts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        return "\n".join(texts)
    raise ValueError("Unsupported file type. Upload a PDF or DOCX resume.")
